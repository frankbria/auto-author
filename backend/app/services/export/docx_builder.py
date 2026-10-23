"""The Word (DOCX) builder (python-docx) (#988)."""
import io
from typing import Dict, List, Optional, BinaryIO

from app.services import inline_marks
from app.services.export.errors import ExportUnavailableError
from app.services.export_templates import PAGE_SIZES_INCHES

# Optional export dependencies are guarded so the service (and the rest of the
# API) still imports if a library is missing — only the affected format fails,
# and it fails loudly with a clear message rather than at import time.
try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


def _add_page_number_field(paragraph) -> None:
    """Append a Word PAGE field to a paragraph (live page number in footer)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


class DocxMixin:
    def _apply_docx_template(self, doc, template: Dict, book_data: Dict) -> None:
        """Apply page size, margins, body font/spacing, and running header/footer."""
        section = doc.sections[0]
        w_in, h_in = PAGE_SIZES_INCHES[template["page_size"]]
        section.page_width = Inches(w_in)
        section.page_height = Inches(h_in)
        m = template["margins"]
        section.top_margin = Inches(m["top"])
        section.bottom_margin = Inches(m["bottom"])
        section.left_margin = Inches(m["inside"])
        section.right_margin = Inches(m["outside"])

        # Body font + line spacing on the Normal style, and the template font on
        # the heading styles so titles/TOC/chapter headings inherit it too.
        font_name = template["font"]["docx_font"]
        normal = doc.styles["Normal"]
        normal.font.name = font_name
        normal.font.size = Pt(template["font"]["size"])
        normal.paragraph_format.line_spacing = template["line_height"]
        for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
            try:
                doc.styles[style_name].font.name = font_name
            except KeyError:
                pass  # style not present in this document's base template

        # Running header (book/author) — page number lives in the footer. A right
        # tab stop at the content width pins the right-hand value to the margin.
        left = self._resolve_tokens(template.get("header", {}).get("left", ""), book_data)
        right = self._resolve_tokens(template.get("header", {}).get("right", ""), book_data)
        if left or right:
            hp = section.header.paragraphs[0]
            if right:
                content_width = Inches(w_in - m["inside"] - m["outside"])
                hp.paragraph_format.tab_stops.add_tab_stop(
                    content_width, WD_TAB_ALIGNMENT.RIGHT
                )
                hp.text = f"{left}\t{right}"
            else:
                hp.text = left
        if template.get("footer", {}).get("center"):
            fp = section.footer.paragraphs[0]
            fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_page_number_field(fp)

    def _build_docx(
        self,
        book_data: Dict,
        chapters: List[Dict],
        output_stream: Optional[BinaryIO] = None,
        template: Optional[Dict] = None,
    ) -> bytes:
        if not DOCX_AVAILABLE:
            raise ExportUnavailableError(
                "DOCX export unavailable: python-docx is not installed"
            )

        # Create document
        doc = Document()

        if template:
            self._apply_docx_template(doc, template, book_data)

        # Title page
        title = doc.add_heading(book_data.get('title', 'Untitled'), 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if book_data.get('subtitle'):
            subtitle = doc.add_paragraph(book_data['subtitle'])
            subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
            subtitle.runs[0].font.size = Pt(18)
            subtitle.runs[0].font.color.rgb = RGBColor(102, 102, 102)

        doc.add_paragraph()  # Empty line

        # Author
        if book_data.get('author_name'):
            author = doc.add_paragraph(f"by {book_data['author_name']}")
            author.alignment = WD_ALIGN_PARAGRAPH.CENTER
            author.runs[0].font.size = Pt(14)

        # About the Author (from the owner's profile bio)
        if book_data.get('author_bio'):
            doc.add_paragraph()
            about = doc.add_heading("About the Author", level=2)
            about.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bio_p = doc.add_paragraph(book_data['author_bio'])
            bio_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Metadata
        metadata_parts = []
        if book_data.get('genre'):
            metadata_parts.append(f"Genre: {book_data['genre']}")
        if book_data.get('target_audience'):
            metadata_parts.append(f"Target Audience: {book_data['target_audience']}")

        if metadata_parts:
            doc.add_paragraph()
            metadata = doc.add_paragraph(' • '.join(metadata_parts))
            metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Description
        if book_data.get('description'):
            doc.add_paragraph()
            doc.add_paragraph(book_data['description'])

        # Page break
        doc.add_page_break()

        # Table of Contents
        doc.add_heading('Table of Contents', 1)

        for i, chapter in enumerate(chapters, 1):
            toc_entry = doc.add_paragraph(" ".join(self._contents_entry(chapter, i)))
            toc_entry.paragraph_format.left_indent = Inches(0.25 * chapter.get('level', 1))

        doc.add_page_break()

        # Chapters
        for i, chapter in enumerate(chapters, 1):
            # Chapter heading, one heading level down for each level of nesting
            doc.add_heading(self._chapter_heading(chapter, i), min(chapter.get('level', 1), 9))

            # Chapter description
            if chapter.get('description'):
                desc = doc.add_paragraph(chapter['description'])
                desc.runs[0].italic = True

            # Chapter content
            content = chapter.get('content', '')
            if content:
                # Extract and format content
                formatted_content = self._extract_text_formatting(content)

                for para in formatted_content:
                    level = {'heading1': 1, 'heading2': 2, 'heading3': 3}.get(para['style'])
                    if level:
                        inline_marks.add_runs(doc.add_heading('', level), para['text'])
                    else:
                        p = doc.add_paragraph()
                        inline_marks.add_runs(p, para['text'])
                        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif not chapter.get('heading_only'):
                doc.add_paragraph("(No content yet)")
                doc.paragraphs[-1].runs[0].italic = True

            # Add page break after each chapter except the last
            if i < len(chapters) and not chapter.get('heading_only'):
                doc.add_page_break()

        # Save to stream
        if output_stream is None:
            output_stream = io.BytesIO()

        doc.save(output_stream)

        # Get bytes if using BytesIO
        if isinstance(output_stream, io.BytesIO):
            output_stream.seek(0)
            return output_stream.getvalue()

        return b''
