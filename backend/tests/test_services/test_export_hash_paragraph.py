"""A paragraph that begins `# ` is a paragraph in the PDF, DOCX and EPUB (#1026).

The builders are handed chapter text and called a block a heading when it began
with `# `, so a paragraph an author typed that way was set as a heading, with the
hash dropped. Real <h1>-<h3> keep their level.
"""

import zipfile
from io import BytesIO

import pytest
from docx import Document

from tests.test_services.test_export_inline_formatting import _export, _pdf_runs

pytestmark = pytest.mark.asyncio

HASHED = ["# not a heading", "## still not", "### nor this"]
PARAGRAPHS = "".join(f"<p>{t}</p>" for t in HASHED) + "<p>plain body</p>"
HEADINGS = "<h1>One</h1><h2>Two</h2><h3>Three</h3>"


async def test_docx_paragraphs_stay_paragraphs():
    doc = Document(BytesIO(await _export(PARAGRAPHS, "docx")))
    styles = {p.text: p.style.name for p in doc.paragraphs}
    for text in HASHED:
        assert styles[text] == styles["plain body"]


async def test_docx_headings_keep_their_level():
    doc = Document(BytesIO(await _export(HEADINGS, "docx")))
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert (styles["One"], styles["Two"], styles["Three"]) == ("Heading 1", "Heading 2", "Heading 3")


async def _epub_body(content: str) -> str:
    with zipfile.ZipFile(BytesIO(await _export(content, "epub"))) as zf:
        return "".join(zf.read(n).decode() for n in zf.namelist() if n.endswith("html"))


async def test_epub_paragraphs_stay_paragraphs():
    body = await _epub_body(PARAGRAPHS)
    for text in HASHED:
        assert f"<p>{text}</p>" in body


async def test_epub_headings_keep_their_level():
    body = await _epub_body(HEADINGS)
    for tag, text in (("h2", "One"), ("h3", "Two"), ("h4", "Three")):
        assert f"<{tag}>{text}</{tag}>" in body


async def test_pdf_paragraphs_are_set_in_the_body_font():
    runs = await _pdf_runs(PARAGRAPHS)
    font_of = {text: font for font, text in runs}
    for text in HASHED:
        assert font_of[text] == font_of["plain body"]


async def test_pdf_headings_are_not_set_in_the_body_font():
    runs = await _pdf_runs(HEADINGS + "<p>plain body</p>")
    font_of = {text: font for font, text in runs}
    assert all(font_of[t] != font_of["plain body"] for t in ("One", "Two", "Three"))
