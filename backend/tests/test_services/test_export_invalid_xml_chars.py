"""Code points XML 1.0 forbids in chapter text, links, titles and author (#1199).

U+FFFE and a C0 control made DOCX raise `ValueError: All strings must be XML
compatible` (a 500) and left the EPUB not well-formed. They are dropped once,
where ExportService takes the book, and the text around them survives.
"""

import zipfile
from io import BytesIO
from xml.dom import minidom

import pytest
from docx import Document

from app.services.export_service import export_service

pytestmark = pytest.mark.asyncio

BAD = "￾\x01\x0c\ud800"
CONTENT = f'<p>before{BAD}after <a href="https://example.com/a{BAD}b">li{BAD}nk</a></p>'


def _book() -> dict:
    return {
        "title": f"Ti{BAD}tle",
        "author_name": f"Au{BAD}thor",
        "table_of_contents": {"chapters": [
            {"id": "c1", "title": f"Chap{BAD}ter", "content": CONTENT, "subchapters": []}
        ]},
    }


async def test_docx():
    doc = Document(BytesIO(await export_service.export_book(_book(), format="docx")))
    assert any("beforeafter link" in p.text for p in doc.paragraphs)
    assert any(p.text == "Chapter" or p.text.endswith("Chapter") for p in doc.paragraphs)
    [link] = [h for p in doc.paragraphs for h in p.hyperlinks]
    assert link.url == "https://example.com/ab"
    assert doc.core_properties.title in ("", "Title")


async def test_epub_is_well_formed():
    with zipfile.ZipFile(BytesIO(await export_service.export_book(_book(), format="epub"))) as zf:
        xml = [zf.read(n).decode() for n in zf.namelist() if n.endswith(("html", "opf", "ncx"))]
    for doc in xml:
        minidom.parseString(doc.encode())
    joined = "".join(xml)
    assert "beforeafter" in joined and "Chapter" in joined and "https://example.com/ab" in joined


async def test_pdf():
    pdf = await export_service.export_book(_book(), format="pdf")
    assert pdf.startswith(b"%PDF") and b"https://example.com/ab" in pdf
