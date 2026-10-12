"""Exports number chapters by their place in the outline (#805).

Every format numbered the flattened list, so a book with two chapters and two
subchapters under the first came out as "Chapter 1" to "Chapter 4", and the
subchapters looked like chapters. They are now Chapter 1, 1.1, 1.2, Chapter 2,
in the body and in each format's table of contents.
"""

import re
import zipfile
from io import BytesIO

import pytest
from docx import Document

from app.services.export_service import export_service

pytestmark = pytest.mark.asyncio

HEADINGS = ["Chapter 1: Planning", "1.1 Outlines", "1.2 Research", "Chapter 2: Writing"]
CONTENTS = ["1. Planning", "1.1 Outlines", "1.2 Research", "2. Writing"]


def _chapter(title, *subchapters, content=None):
    return {
        "id": title.lower(),
        "title": title,
        "content": f"<p>About {title.lower()}.</p>" if content is None else content,
        "subchapters": list(subchapters),
    }


def _book(*chapters):
    return {"title": "Numbering", "table_of_contents": {"chapters": list(chapters)}}


def two_by_two():
    """Two chapters, two subchapters under the first."""
    return _book(
        _chapter("Planning", _chapter("Outlines"), _chapter("Research")), _chapter("Writing")
    )


async def _docx(book, **options):
    return Document(BytesIO(await export_service.export_book(book, format="docx", **options)))


def _docx_headings(doc):
    """(style, text) of every heading after the table of contents."""
    found = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("Heading")]
    return found[[text for _, text in found].index("Table of Contents") + 1 :]


async def _epub_files(book):
    with zipfile.ZipFile(BytesIO(await export_service.export_book(book, format="epub"))) as zf:
        return {name: zf.read(name).decode() for name in zf.namelist() if name.endswith("html")}


def _text(markup: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", markup)).strip()


class TestATwoByTwoBookIsNumberedByItsOutline:
    async def test_docx_headings_and_their_levels(self):
        doc = await _docx(two_by_two())

        assert _docx_headings(doc) == [
            ("Heading 1", "Chapter 1: Planning"),
            ("Heading 2", "1.1 Outlines"),
            ("Heading 2", "1.2 Research"),
            ("Heading 1", "Chapter 2: Writing"),
        ]

    async def test_docx_table_of_contents(self):
        texts = [p.text for p in (await _docx(two_by_two())).paragraphs]

        start = texts.index("Table of Contents") + 1
        assert texts[start : start + 4] == CONTENTS

    async def test_pdf_headings_and_table_of_contents(self, monkeypatch):
        from reportlab import rl_config

        monkeypatch.setattr(rl_config, "pageCompression", 0)
        monkeypatch.setattr(rl_config, "useA85", 0)

        pdf = await export_service.export_book(two_by_two(), format="pdf")

        # ReportLab writes text as (..) Tj fragments; compare with whitespace removed.
        text = re.sub(rb"\s+", b"", b"".join(re.findall(rb"\(([^)]*)\) Tj", pdf))).decode()
        compact = [re.sub(r"\s+", "", line) for line in CONTENTS + HEADINGS]
        positions = [text.find(line, 0 if i < 4 else text.find(compact[3])) for i, line in enumerate(compact)]
        assert -1 not in positions, list(zip(compact, positions))
        assert positions == sorted(positions)
        assert "Chapter3" not in text and "Chapter4" not in text

    async def test_pdf_sets_a_subchapter_heading_smaller_than_a_chapter_heading(self, monkeypatch):
        from reportlab import rl_config

        monkeypatch.setattr(rl_config, "pageCompression", 0)
        monkeypatch.setattr(rl_config, "useA85", 0)

        pdf = await export_service.export_book(two_by_two(), format="pdf")

        # Each text fragment with the font size in force when it was drawn.
        sizes, size = {}, None
        for set_size, text in re.findall(rb"/F\d+ ([\d.]+) Tf|\(([^)]*)\) Tj", pdf):
            if set_size:
                size = float(set_size)
            else:
                sizes[text.decode()] = size
        assert sizes["1.1 Outlines"] == sizes["1.2 Research"]
        assert sizes["1.1 Outlines"] < sizes["Chapter 1: Planning"] == sizes["Chapter 2: Writing"]

    async def test_epub_headings(self):
        files = await _epub_files(two_by_two())

        chapters = [files[name] for name in sorted(files) if "chap_" in name]
        assert [re.search(r"<(h\d)[^>]*>(.*?)</\1>", c, flags=re.S).groups() for c in chapters] == [
            ("h1", "Chapter 1: Planning"),
            ("h2", "1.1 Outlines"),
            ("h2", "1.2 Research"),
            ("h1", "Chapter 2: Writing"),
        ]

    async def test_epub_navigation_nests_the_subchapters_under_their_chapter(self):
        files = await _epub_files(two_by_two())

        nav = next(markup for name, markup in files.items() if "nav" in name)
        planning = nav[nav.index("Chapter 1: Planning") : nav.index("Chapter 2: Writing")]
        assert _text(planning) == "Chapter 1: Planning 1.1 Outlines 1.2 Research"
        # The two subchapters sit in a list of their own, inside the chapter's entry.
        assert planning.count("<ol") == 1 and planning.count("</ol>") == 1

    async def test_markdown_headings(self):
        markdown = (await export_service.export_book(two_by_two(), format="markdown")).decode()

        assert [line for line in markdown.splitlines() if line.startswith("##")] == [
            "## Chapter 1: Planning",
            "### 1.1 Outlines",
            "### 1.2 Research",
            "## Chapter 2: Writing",
        ]


class TestNumberingFollowsWhatIsExported:
    async def test_a_chapter_with_no_text_of_its_own_still_heads_its_subchapters(self):
        """Empty chapters are left out by default, but a chapter that only groups
        subchapters is the usual shape of a book, and its subchapters need it."""
        book = _book(
            _chapter("Planning", _chapter("Outlines"), _chapter("Research"), content=""),
            _chapter("Writing"),
        )

        assert [text for _, text in _docx_headings(await _docx(book))] == HEADINGS

    async def test_a_heading_that_only_introduces_subchapters_has_no_placeholder_and_no_page_of_its_own(
        self, monkeypatch
    ):
        """The "(No content yet)" line is for a chapter the author asked to see
        empty, not for one whose text is all in its subchapters."""
        from reportlab import rl_config

        monkeypatch.setattr(rl_config, "pageCompression", 0)
        monkeypatch.setattr(rl_config, "useA85", 0)
        book = _book(_chapter("Planning", _chapter("Outlines"), content=""), _chapter("Writing"))

        doc = await _docx(book)
        after_toc = [p for p in doc.paragraphs][[p.text for p in doc.paragraphs].index("Chapter 1: Planning") :]
        assert [p.text for p in after_toc[:2]] == ["Chapter 1: Planning", "1.1 Outlines"]
        epub = "".join((await _epub_files(book)).values())
        markdown = (await export_service.export_book(book, format="markdown")).decode()
        pdf = await export_service.export_book(book, format="pdf")
        assert "No content yet" not in "".join(p.text for p in doc.paragraphs) + epub + markdown
        assert b"No content yet" not in pdf
        # Title page, contents, then Chapter 1 with 1.1 under it, then Chapter 2.
        assert len(re.findall(rb"/Type /Page\b", pdf)) == 4

    async def test_an_empty_chapter_that_was_asked_for_keeps_its_placeholder(self):
        book = _book(_chapter("Planning", content=""), _chapter("Writing"))

        doc = await _docx(book, include_empty_chapters=True)

        assert "(No content yet)" in [p.text for p in doc.paragraphs]

    async def test_an_empty_chapter_with_nothing_under_it_is_left_out_and_not_counted(self):
        book = _book(
            _chapter("Planning", _chapter("Outlines", content=""), _chapter("Research")),
            _chapter("Notes", _chapter("Scraps", content=""), content=""),
            _chapter("Writing"),
        )

        assert [text for _, text in _docx_headings(await _docx(book))] == [
            "Chapter 1: Planning",
            "1.1 Research",
            "Chapter 2: Writing",
        ]

    async def test_empty_chapters_are_numbered_when_they_are_asked_for(self):
        book = _book(
            _chapter("Planning", _chapter("Outlines", content="")), _chapter("Writing", content="")
        )

        doc = await _docx(book, include_empty_chapters=True)

        assert [text for _, text in _docx_headings(doc)] == [
            "Chapter 1: Planning",
            "1.1 Outlines",
            "Chapter 2: Writing",
        ]

    async def test_a_third_level_gets_a_third_number(self):
        book = _book(_chapter("Planning", _chapter("Outlines", _chapter("Scenes"))))

        assert _docx_headings(await _docx(book)) == [
            ("Heading 1", "Chapter 1: Planning"),
            ("Heading 2", "1.1 Outlines"),
            ("Heading 3", "1.1.1 Scenes"),
        ]
