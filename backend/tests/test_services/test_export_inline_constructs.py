"""Links, inline code, strike-through, superscript and subscript in the PDF, DOCX
and EPUB (#1190), and `# ` at the start of a list item or a bare block.

The builders were handed Markdown's `[text](url)`, backticks and `~~`, and
printed them; superscripts lost their position (`x2`). Text that only looks
like markup (`[not a link]`, `__init__`) must still print as typed, and the
Markdown export keeps its syntax.
"""

from io import BytesIO

import pytest
from docx import Document

from tests.test_services.test_export_inline_formatting import (
    _docx_paragraph, _epub_chapter, _export, _pdf_runs,
)

pytestmark = pytest.mark.asyncio

LINK = '<p>see <a href="https://example.com/a?b=1&amp;c=2">the <strong>site</strong></a> now</p>'
CODE = "<p>call <code>a_b *c*</code> here</p>"
STRIKE = "<p>old <s>gone</s> and <del>deleted</del> text</p>"
SCRIPTS = "<p>E = mc<sup>2</sup> and H<sub>2</sub>O</p>"
LOOKALIKES = "<p>[not a link] and [a](b) and `tick` and ~~tilde~~ and x^2</p>"
UNSAFE = '<p>go <a href="javascript:alert(1)">there</a> or <a>nowhere</a></p>'


class TestLinks:
    async def test_docx_is_a_hyperlink(self):
        doc = Document(BytesIO(await _export(LINK, "docx")))
        paragraph = next(p for p in doc.paragraphs if p.text == "see the site now")
        [link] = paragraph.hyperlinks
        assert link.url == "https://example.com/a?b=1&c=2"
        assert link.text == "the site"

    async def test_epub_is_an_anchor(self):
        chapter = await _epub_chapter(LINK)
        assert '<a href="https://example.com/a?b=1&amp;c=2">the <strong>site</strong></a>' in chapter

    async def test_pdf_is_a_link_annotation(self):
        pdf = await _export(LINK, "pdf")
        assert b"/URI (https://example.com/a?b=1&c=2)" in pdf
        drawn = "".join(text for _, text in await _pdf_runs(LINK))
        assert "see the site now" in drawn and "](" not in drawn

    @pytest.mark.parametrize("fmt", ["docx", "epub", "pdf"])
    async def test_only_web_and_mail_addresses_become_links(self, fmt):
        data = await _export(UNSAFE, fmt)
        if fmt == "docx":
            doc = Document(BytesIO(data))
            assert not any(p.hyperlinks for p in doc.paragraphs)
            assert any(p.text == "go there or nowhere" for p in doc.paragraphs)
        elif fmt == "epub":
            chapter = await _epub_chapter(UNSAFE)
            assert "javascript" not in chapter and "<p>go there or nowhere</p>" in chapter
        else:
            assert b"javascript" not in data
            assert "go there or nowhere" in "".join(t for _, t in await _pdf_runs(UNSAFE))


class TestInlineCode:
    async def test_docx_is_monospace(self):
        paragraph = await _docx_paragraph(CODE, "call a_b *c* here")
        assert [(r.text, r.font.name) for r in paragraph.runs if r.text == "a_b *c*"] == [("a_b *c*", "Courier New")]

    async def test_epub(self):
        assert "call <code>a_b *c*</code> here" in await _epub_chapter(CODE)

    async def test_pdf(self):
        runs = await _pdf_runs(CODE)
        assert any(text == "a_b *c*" and "Courier" in font for font, text in runs), runs
        assert not any("`" in text for _, text in runs)


class TestStrikeThrough:
    async def test_docx(self):
        paragraph = await _docx_paragraph(STRIKE, "old gone and deleted text")
        assert {r.text for r in paragraph.runs if r.font.strike} == {"gone", "deleted"}

    async def test_epub(self):
        chapter = await _epub_chapter(STRIKE)
        assert "<s>gone</s>" in chapter and "<s>deleted</s>" in chapter

    async def test_pdf(self):
        drawn = "".join(text for _, text in await _pdf_runs(STRIKE))
        assert "old gone and deleted text" in drawn and "~" not in drawn


class TestSuperscriptAndSubscript:
    async def test_docx(self):
        paragraph = await _docx_paragraph(SCRIPTS, "E = mc2 and H2O")
        assert [r.text for r in paragraph.runs if r.font.superscript] == ["2"]
        assert [r.text for r in paragraph.runs if r.font.subscript] == ["2"]

    async def test_epub(self):
        chapter = await _epub_chapter(SCRIPTS)
        assert "mc<sup>2</sup>" in chapter and "H<sub>2</sub>O" in chapter

    async def test_pdf(self):
        drawn = "".join(text for _, text in await _pdf_runs(SCRIPTS))
        assert "E = mc2 and H2O" in drawn


class TestLookalikesAreAsTyped:
    async def test_docx(self):
        paragraph = await _docx_paragraph(LOOKALIKES, "[not a link] and [a](b) and `tick` and ~~tilde~~ and x^2")
        assert not paragraph.hyperlinks
        assert not any(r.bold or r.italic or r.font.strike or r.font.name for r in paragraph.runs)

    async def test_epub(self):
        assert "<p>[not a link] and [a](b) and `tick` and ~~tilde~~ and x^2</p>" in await _epub_chapter(LOOKALIKES)

    async def test_pdf(self):
        drawn = "".join(text for _, text in await _pdf_runs(LOOKALIKES)).replace("\\", "")  # PDF string escapes
        assert "[not a link] and [a](b) and `tick` and ~~tilde~~ and x^2" in drawn


async def test_the_markdown_export_keeps_its_syntax():
    markdown = (await _export(LINK + CODE + STRIKE, "markdown")).decode()
    assert "[the **site**](https://example.com/a?b=1&c=2)" in markdown
    assert "`a_b *c*`" in markdown and "~~gone~~" in markdown


HASHED_BLOCKS = "<ul><li># in a list</li></ul><div># in a div</div><div>## second div</div># bare text"
EXPECTED = ["# in a div", "## second div", "# bare text"]


class TestHashAtTheStartOfABlockIsText:
    async def test_docx(self):
        doc = Document(BytesIO(await _export(HASHED_BLOCKS, "docx")))
        styles = {p.text: p.style.name for p in doc.paragraphs}
        for text in EXPECTED + ["* # in a list"]:
            assert styles[text] == "Normal", styles

    async def test_epub(self):
        chapter = await _epub_chapter(HASHED_BLOCKS)
        for text in EXPECTED + ["* # in a list"]:
            assert f"<p>{text}</p>" in chapter

    async def test_pdf(self):
        runs = await _pdf_runs(HASHED_BLOCKS + "<p>plain body</p>")
        font_of = {text: font for font, text in runs}
        for text in EXPECTED + ["* # in a list"]:
            assert font_of[text] == font_of["plain body"], runs
